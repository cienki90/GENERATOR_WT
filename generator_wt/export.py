"""Zestawienia w Excelu na podstawie szablonów z katalogu szablony/.

rozbudowana.xlsx  - arkusz 'Zał.1': nagłówek (wiersze 1-3) z szablonu, wiersze danych
                    formatowane jak pierwszy wiersz danych szablonu, stopka z liczbą słupów.
uproszczona.xls   - układ jak w szablonie (Załącznik nr 4, wykaz, stopka z podpisami);
                    teksty nagłówka są czytane z szablonu, plik budowany od nowa (xlwt).
"""
from __future__ import annotations

import copy
import re
from pathlib import Path

import xlrd
import xlwt
from openpyxl import load_workbook

from .config import Config
from .model import Wierzcholek

STREFY = {2176: 5, 2177: 6, 2178: 7, 2179: 8}


def _opis_ukladu(tekst: str, epsg: int | None) -> str:
    if epsg in STREFY:
        return re.sub(r"str\.\s*\d", f"str. {STREFY[epsg]}", tekst)
    return tekst


def _rejon(w: Wierzcholek) -> str:
    return w.rejon_skrot or w.rejon or ""


# ------------------------------------------------------------------ rozbudowana

def zapisz_rozbudowane(slupy: list[Wierzcholek], plik: Path, cfg: Config,
                       epsg: int | None = None) -> None:
    wb = load_workbook(cfg.szablon_rozbudowana)
    ws = wb.worksheets[0]
    for inny in wb.worksheets[1:]:  # arkusze robocze szablonu nie trafiają do wyniku
        wb.remove(inny)
    wiersz_wzor = 4
    styl = {c.column: copy.copy(c._style) for c in ws[wiersz_wzor]}
    wys = ws.row_dimensions[wiersz_wzor].height
    stopka_styl = copy.copy(ws.cell(ws.max_row, 1)._style)
    for zakres in list(ws.merged_cells.ranges):
        if zakres.min_row >= wiersz_wzor:
            ws.unmerge_cells(str(zakres))
    ws.delete_rows(wiersz_wzor, ws.max_row - wiersz_wzor + 1)
    ws["F3"] = _opis_ukladu(ws["F3"].value or "", epsg)

    fmt = f"{{:.{cfg.miejsca_po_przecinku}f}}"
    for lp, w in enumerate(slupy, start=1):
        r = wiersz_wzor + lp - 1
        wartosci = [lp, _rejon(w), w.gmina or "", w.miejscowosc or "", w.ulica or "-",
                    fmt.format(w.geo_x), fmt.format(w.geo_y), None, w.stacja_trafo or "",
                    w.rodzaj_slupa or "", cfg.linie_swiatlowodowe, cfg.linie_abonenckie]
        for k, v in enumerate(wartosci, start=1):
            c = ws.cell(r, k, v)
            if k in styl:
                c._style = copy.copy(styl[k])
        ws.row_dimensions[r].height = wys
    r = wiersz_wzor + len(slupy) + 1
    c = ws.cell(r, 1, f"Wykaz dotyczy łącznie {len(slupy)} słupów.")
    c._style = stopka_styl
    ws.print_title_rows = "1:3"
    wb.save(plik)


# ------------------------------------------------------------------ uproszczona

def _teksty_szablonu_xls(plik: Path) -> dict[str, str]:
    domyslne = {
        "zalacznik": "Załącznik nr 4 do Umowy nr ………………………………..",
        "tytul": "Wykaz Słupów Elektroenergetycznych wykorzystywanych do podwieszenia "
                 "Przewodu Telekomunikacyjnego Najemcy",
        "naglowki": ["Nr słupa", "Rejon Energetyczny", "Gmina", "Miejscowość",
                     "Lokalizacja, ulica", "Współrzędne słupa w układzie PUWG 2000 str. 7"],
        "uwaga": "(słupy należy wpisywać w kolejnych wierszach, w razie potrzeby dodać "
                 "wymaganą ilość pozycji) ",
        "strona_l": "WNIOSKODAWCA/NAJEMCA", "strona_p": "WYNAJMUJĄCY",
        "kropki_l": "…………………...…..…………………….", "kropki_p": "……………………………….…………..",
        "podpis": "data i podpis",
    }
    if not plik.exists():
        return domyslne
    s = xlrd.open_workbook(str(plik)).sheet_by_index(0)
    t = dict(domyslne)
    t["zalacznik"] = s.cell_value(0, 4) or t["zalacznik"]
    t["tytul"] = s.cell_value(2, 0) or t["tytul"]
    t["naglowki"] = [s.cell_value(3, c) or d for c, d in enumerate(domyslne["naglowki"])]
    for r in range(4, s.nrows):
        v = str(s.cell_value(r, 0))
        if v.startswith("(słupy"):
            t["uwaga"] = v
        elif v.isupper() and s.cell_value(r, 4):
            t["strona_l"], t["strona_p"] = v, s.cell_value(r, 4)
        elif v.startswith("…"):
            t["kropki_l"], t["kropki_p"] = v, s.cell_value(r, 4) or t["kropki_p"]
    return t


def zapisz_uproszczone(slupy: list[Wierzcholek], plik: Path, cfg: Config,
                       epsg: int | None = None) -> None:
    t = _teksty_szablonu_xls(cfg.szablon_uproszczona)
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("Summary")
    ramka = "borders: left thin, right thin, top thin, bottom thin;"
    sr = "align: horiz center, vert center, wrap on;"
    f10b = "font: name Calibri, height 200, bold on;"
    st_zal = xlwt.easyxf(f10b + "align: horiz center;")
    st_tytul = xlwt.easyxf(f10b + sr)
    st_nagl = xlwt.easyxf("font: name Calibri, height 200;" + sr + ramka)
    st_txt = xlwt.easyxf("font: name Calibri, height 220;" + sr + ramka, num_format_str="@")
    st_num = xlwt.easyxf("font: name Calibri, height 220;" + sr + ramka, num_format_str="0.00")
    st_stopka = xlwt.easyxf(f10b + "align: horiz center;")
    st_zwykly = xlwt.easyxf("font: name Calibri, height 200;" + "align: horiz center;")

    for k, szer in enumerate([9, 18, 18, 19, 20, 14, 15]):
        ws.col(k).width = 256 * szer

    def wys(r, h):
        ws.row(r).height_mismatch = True
        ws.row(r).height = h

    ws.write_merge(0, 0, 4, 6, t["zalacznik"], st_zal)
    ws.write_merge(2, 2, 0, 6, t["tytul"], st_tytul)
    wys(2, 525)
    for k, n in enumerate(t["naglowki"][:5]):
        ws.write(3, k, n, st_nagl)
    ws.write_merge(3, 3, 5, 6, _opis_ukladu(t["naglowki"][5], epsg), st_nagl)
    wys(3, 660)

    for i, w in enumerate(slupy):
        r = 4 + i
        ws.write(r, 0, f"  {w.etykieta}", st_txt)
        ws.write(r, 1, _rejon(w), st_txt)
        ws.write(r, 2, w.gmina or "", st_txt)
        ws.write(r, 3, w.miejscowosc or "", st_txt)
        ws.write(r, 4, w.ulica or "-", st_txt)
        ws.write(r, 5, round(w.geo_x, cfg.miejsca_po_przecinku), st_num)
        ws.write(r, 6, round(w.geo_y, cfg.miejsca_po_przecinku), st_num)
        wys(r, 300)

    r = 4 + len(slupy)
    ws.write_merge(r, r, 0, 6, t["uwaga"], st_zwykly)
    ws.write_merge(r + 2, r + 2, 0, 6, f"Wykaz dotyczy łącznie {len(slupy)} szt. słupów.",
                   st_stopka)
    ws.write_merge(r + 5, r + 5, 0, 3, t["strona_l"], st_stopka)
    ws.write_merge(r + 5, r + 5, 4, 6, t["strona_p"], st_stopka)
    ws.write_merge(r + 8, r + 8, 0, 3, t["kropki_l"], st_zwykly)
    ws.write_merge(r + 8, r + 8, 4, 6, t["kropki_p"], st_zwykly)
    ws.write_merge(r + 9, r + 9, 0, 3, t["podpis"], st_zwykly)
    ws.write_merge(r + 9, r + 9, 4, 6, t["podpis"], st_zwykly)
    ws.set_panes_frozen(True)
    ws.set_horz_split_pos(4)
    wb.save(str(plik))
