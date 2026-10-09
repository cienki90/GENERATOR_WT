"""Wypełnianie pisma 'Zapytanie o możliwość dostępu do słupów' (szablon docx).

Szablon: szablony/pismo_warunki.docx - program podmienia:
- dane operatora (pola 'Etykieta<TAB>wartość'),
- okres 'OD: ... DO: ...',
- adres słupa początkowego i końcowego, łączną ilość słupów,
- tabelę 'Załącznik 1.1 Wykaz słupów' (wiersze generowane wg numeracji).
Formatowanie szablonu zostaje zachowane.
"""
from __future__ import annotations

import copy
from datetime import date
from pathlib import Path

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

from .model import Wierzcholek

# etykieta w piśmie -> kolumna w arkuszu Operatorzy
POLA_OPERATORA = {
    "Numer Umowy Ramowej": "Numer Umowy Ramowej",
    "Pełna nazwa": "Pełna nazwa",
    "Adres (siedziba)": "Adres (siedziba)",
    "Kod pocztowy": "Kod pocztowy",
    "NIP": "NIP",
    "Regon": "Regon",
    "Numer wpisu do RPT": "Numer wpisu do RPT",
}


def _ustaw_tekst(p: Paragraph, tekst: str) -> None:
    """Zastępuje tekst akapitu, zachowując formatowanie pierwszego fragmentu."""
    if not p.runs:
        p.add_run(tekst)
        return
    p.runs[0].text = tekst
    for r in p.runs[1:]:
        r.text = ""


def _ustaw_wartosc_po_tabulatorze(p: Paragraph, wartosc: str) -> None:
    """Dla akapitów 'Etykieta<TAB>wartość' podmienia tylko część po ostatnim tabulatorze."""
    runs = p.runs
    ost_tab = max((i for i, r in enumerate(runs) if "\t" in r.text), default=None)
    if ost_tab is None:
        _ustaw_tekst(p, wartosc)
        return
    # tabulator mógł być połączony z tekstem w jednym fragmencie
    r = runs[ost_tab]
    po_tab = r.text[r.text.rfind("\t") + 1:]
    if po_tab.strip():  # wartość sklejona z tabulatorem w jednym fragmencie
        r.text = r.text[: r.text.rfind("\t") + 1]
    if ost_tab + 1 < len(runs):
        runs[ost_tab + 1].text = wartosc
        for x in runs[ost_tab + 2:]:
            x.text = ""
    else:
        r.text += wartosc


def _akapit(paragrafy: list[Paragraph], poczatek: str, od: int = 0) -> int | None:
    for i in range(od, len(paragrafy)):
        if paragrafy[i].text.strip().startswith(poczatek):
            return i
    return None


def _nastepny_niepusty(paragrafy: list[Paragraph], od: int) -> int | None:
    for i in range(od, len(paragrafy)):
        if paragrafy[i].text.strip():
            return i
    return None


def adres_slupa(w: Wierzcholek) -> str:
    adres = w.miejscowosc or ""
    if w.ulica:
        adres += f" ul. {w.ulica}"
    poczta = " ".join(x for x in (w.kod, w.gmina) if x)
    return f"{adres}, {poczta}" if poczta else adres


def _ustaw_komorke(komorka, tekst: str) -> None:
    p = komorka.paragraphs[0]
    _ustaw_tekst(p, tekst)
    for nadmiar in komorka.paragraphs[1:]:
        nadmiar._element.getparent().remove(nadmiar._element)


def _tabela_slupow(d: docx.Document) -> Table:
    for t in d.tables:
        if t.rows and t.rows[0].cells[0].text.strip().startswith("L.p"):
            return t
    raise ValueError("Nie znaleziono tabeli wykazu słupów (nagłówek 'L.p.').")


def wypelnij_tabele(t: Table, slupy: list[Wierzcholek], cfg=None) -> None:
    wgs = bool(cfg and cfg.wgs84)
    m2, mw = (cfg.miejsca_po_przecinku, cfg.miejsca_wgs84) if cfg else (2, 7)
    wzor = copy.deepcopy(t.rows[1]._tr)
    for r in list(t.rows)[1:]:
        t._tbl.remove(r._tr)
    for lp, w in enumerate(slupy, start=1):
        tr = copy.deepcopy(wzor)
        t._tbl.append(tr)
        wiersz = t.rows[-1]
        wartosci = [str(lp), w.rejon or "", w.gmina or "", w.miejscowosc or "",
                    w.ulica or "-"] + [f"{v:.{mw if wgs else m2}f}"
                                       for v in w.wspolrzedne(wgs, m2, mw)]
        for kom, txt in zip(wiersz.cells, wartosci):
            _ustaw_komorke(kom, txt)


def generuj_pismo(szablon: Path, wyjscie: Path, slupy: list[Wierzcholek],
                  operator: dict[str, str] | None, data_od: str | None = None,
                  data_do: str = "-", cfg=None) -> None:
    d = docx.Document(str(szablon))
    par = d.paragraphs

    if operator:
        for etykieta, kolumna in POLA_OPERATORA.items():
            i = _akapit(par, etykieta)
            if i is not None and operator.get(kolumna):
                _ustaw_wartosc_po_tabulatorze(par[i], operator[kolumna])
        i = _akapit(par, "Dane kontaktowe")
        if i is not None and operator.get("Dane kontaktowe"):
            j = _nastepny_niepusty(par, i + 1)
            if j is not None:
                _ustaw_tekst(par[j], operator["Dane kontaktowe"])

    i = _akapit(par, "OD:")
    if i is not None:
        _ustaw_tekst(par[i], f"OD: {data_od or date.today().strftime('%d.%m.%Y')} DO: {data_do}")

    if slupy:
        for etykieta, w in (("Słup początkowy", slupy[0]), ("Słup końcowy", slupy[-1])):
            i = _akapit(par, etykieta)
            j = _akapit(par, "Adres", i + 1) if i is not None else None
            k = _nastepny_niepusty(par, j + 1) if j is not None else None
            if k is not None:
                _ustaw_tekst(par[k], adres_slupa(w))

    i = _akapit(par, "Łączna ilość Słupów")
    if i is not None:
        p = par[i]
        if len(p.runs) >= 2:
            p.runs[0].text = "Łączna ilość Słupów elektroenergetycznych: "
            p.runs[1].text = str(len(slupy))
            for r in p.runs[2:]:
                r.text = ""
        else:
            _ustaw_tekst(p, f"Łączna ilość Słupów elektroenergetycznych: {len(slupy)}")

    wypelnij_tabele(_tabela_slupow(d), slupy, cfg)
    d.save(str(wyjscie))
