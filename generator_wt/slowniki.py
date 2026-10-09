"""Słowniki z Excela: operatorzy i rejony energetyczne (dane/slowniki.xlsx).

Arkusz "Operatorzy" - kolumny odpowiadają polom pisma (patrz KOLUMNY_OPERATORA).
Arkusz "Rejony"     - Nazwa (do pisma), Nazwa skrócona (do zestawień Excel),
                      Gminy (lista gmin rozdzielona średnikami), pozostałe dowolne.
Pierwszy wiersz każdego arkusza to nagłówki. Kolumny można dopisywać.
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

ARKUSZ_OPERATORZY = "Operatorzy"
ARKUSZ_REJONY = "Rejony"

KOLUMNY_OPERATORA = ["Nazwa", "Numer Umowy Ramowej", "Pełna nazwa", "Adres (siedziba)",
                     "Kod pocztowy", "NIP", "Regon", "Numer wpisu do RPT",
                     "Dane kontaktowe"]
KOLUMNY_REJONU = ["Nazwa", "Nazwa skrócona", "Gminy", "Uwagi"]


def wczytaj_liste(plik: Path, arkusz: str) -> list[dict[str, str]]:
    wb = load_workbook(plik, read_only=True, data_only=True)
    if arkusz not in wb.sheetnames:
        raise ValueError(f"Brak arkusza '{arkusz}' w pliku {plik}")
    wiersze = list(wb[arkusz].iter_rows(values_only=True))
    wb.close()
    if not wiersze:
        return []
    naglowki = [str(n).strip() if n is not None else f"Kol{i}" for i, n in enumerate(wiersze[0])]
    wynik = []
    for w in wiersze[1:]:
        if not any(v not in (None, "") for v in w):
            continue
        wynik.append({n: ("" if v is None else str(v).strip()) for n, v in zip(naglowki, w)})
    return wynik


def wybierz(lista: list[dict[str, str]], tytul: str, domyslny: int | None = None
            ) -> dict[str, str] | None:
    """Wybór pozycji z listy w konsoli (docelowo: lista rozwijana w GUI)."""
    if not lista:
        print(f"Lista '{tytul}' jest pusta - uzupełnij plik słowników.")
        return None
    print(f"\n{tytul}:")
    for i, poz in enumerate(lista, start=1):
        print(f"  {i:>3}. {poz.get('Nazwa', '')}")
    podp = f" (Enter = {domyslny})" if domyslny else " (Enter = pomiń)"
    while True:
        odp = input(f"Wybierz numer 1-{len(lista)}{podp}: ").strip()
        if not odp:
            return lista[domyslny - 1] if domyslny else None
        if odp.isdigit() and 1 <= int(odp) <= len(lista):
            return lista[int(odp) - 1]
        print("Nieprawidłowy wybór.")


def _norm(t: str) -> str:
    return re.sub(r"\s+", " ", t.strip().casefold())


def rejon_dla_gminy(rejony: list[dict[str, str]], gmina: str | None
                    ) -> dict[str, str] | None:
    """Rejon energetyczny, w którego kolumnie 'Gminy' występuje dana gmina."""
    if not gmina:
        return None
    g = _norm(gmina)
    for r in rejony:
        gminy = {_norm(x) for x in re.split(r"[;,\n]", r.get("Gminy", "")) if x.strip()}
        if g in gminy:
            return r
    return None


def utworz_szablon(plik: Path) -> None:
    """Tworzy plik słowników z przykładowymi wpisami (na podstawie pisma)."""
    plik.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = ARKUSZ_OPERATORZY
    ws.append(KOLUMNY_OPERATORA)
    ws.append(["InterWan", "1441OW2018", "InterWan Sp z.o.o.", "ul. Lazurowa 1",
               "05-331 Dębę Wielkie", "8222349374", "361015909", "5522",
               "Marcin Macko tel. 502 761 590 marcin.macko@mmui.pl"])
    ws2 = wb.create_sheet(ARKUSZ_REJONY)
    ws2.append(KOLUMNY_REJONU)
    ws2.append(["Rejon Energetyczny Mińsk Mazowiecki", "Mińsk Mazowiecki",
                "Mińsk Mazowiecki; Mińsk Mazowiecki (miasto); Dębe Wielkie; Jakubów", ""])
    ws2.append(["Rejon Energetyczny Otwock", "OTWOCK", "Otwock; Józefów; Karczew; Celestynów; Halinów", ""])
    for arkusz in wb.worksheets:
        for c in arkusz[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDEBF7")
        for kol in arkusz.columns:
            arkusz.column_dimensions[kol[0].column_letter].width = 30
        arkusz.freeze_panes = "A2"
    wb.save(plik)


def nazwa_skrocona(rejon: dict[str, str] | None) -> str | None:
    if not rejon:
        return None
    return rejon.get("Nazwa skrócona") or rejon.get("Nazwa")


def zapisz_operatora(plik: Path, wpis: dict[str, str]) -> str:
    """Dopisuje operatora do arkusza 'Operatorzy' albo aktualizuje wiersz z tym samym NIP.
    Brakujące kolumny są dopisywane na końcu nagłówka. Zwraca 'dodano' / 'zaktualizowano'."""
    wb = load_workbook(plik)
    ws = wb[ARKUSZ_OPERATORZY]
    nagl = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    for k in wpis:
        if k.startswith("_") or k in nagl:
            continue
        nagl.append(k)
        c = ws.cell(1, len(nagl), k)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
        ws.column_dimensions[c.column_letter].width = 30
    kol = {n: i + 1 for i, n in enumerate(nagl) if n}
    nip = re.sub(r"\D", "", wpis.get("NIP", ""))
    wiersz, stan = None, "dodano"
    if nip and "NIP" in kol:
        for r in range(2, ws.max_row + 1):
            if re.sub(r"\D", "", str(ws.cell(r, kol["NIP"]).value or "")) == nip:
                wiersz, stan = r, "zaktualizowano"
                break
    if wiersz is None:
        wiersz = ws.max_row + 1
        while wiersz > 2 and not any(ws.cell(wiersz - 1, k).value for k in kol.values()):
            wiersz -= 1
    for k, v in wpis.items():
        if k in kol and (v or stan == "dodano"):
            ws.cell(wiersz, kol[k], v)
    wb.save(plik)
    return stan
