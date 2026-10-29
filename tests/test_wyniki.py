"""Pliki wyników: DXF z numeracją i arkuszami, PDF, pismo, zestawienia, wiersz poleceń."""
import docx
import ezdxf
import pytest
import xlrd
from openpyxl import load_workbook

from generator_wt import slowniki
from generator_wt.main import main as cli
from tests.conftest import nowy_projekt


@pytest.fixture
def operator(dxf):
    from generator_wt.config import Config
    return slowniki.wczytaj_liste(Config().plik_slownikow, slowniki.ARKUSZ_OPERATORZY)[0]


def test_dxf_numeracja_i_arkusze(dxf, operator):
    p = nowy_projekt(dxf, arkusze=True)
    cel = dxf.with_name("projekt_numeracja.dxf")
    p.zapisz_dxf(cel, True, operator, "Test", "10.2026")
    doc = ezdxf.readfile(cel)
    msp = doc.modelspace()
    warstwa = p.cfg.warstwa_numeracji
    teksty = msp.query(f'MTEXT[layer=="{warstwa}"]')
    assert len(msp.query(f'LEADER[layer=="{warstwa}"]')) == len(p.plan)
    assert sorted(t.text for t in teksty) == sorted(w.etykieta for w in p.plan)
    uklady = sorted((n for n in doc.layouts.names() if n.isdigit()), key=int)
    assert uklady == [str(i) for i in range(1, len(p.arkusze) + 1)]
    # oryginał bez zmian
    assert len(ezdxf.readfile(dxf).modelspace().query(f'MTEXT[layer=="{warstwa}"]')) == 0


def test_dxf_ponowny_zapis_zastepuje_numeracje(dxf):
    p = nowy_projekt(dxf)
    cel = dxf.with_name("projekt_numeracja.dxf")
    p.zapisz_dxf(cel)
    p.zapisz_dxf(cel)  # drugi zapis na tym samym dokumencie - bez dublowania
    warstwa = p.cfg.warstwa_numeracji
    assert len(ezdxf.readfile(cel).modelspace().query(f'LEADER[layer=="{warstwa}"]')) \
        == len(p.plan)


def test_pdf_arkuszy(dxf, operator):
    import pymupdf
    p = nowy_projekt(dxf, arkusze=True)
    p.zapisz_dxf(dxf.with_name("projekt_numeracja.dxf"), True, operator)
    cel = dxf.with_name("projekt_numeracja.pdf")
    assert p.zapisz_numeracje_pdf(cel, log=lambda t: None) == len(p.arkusze)
    with pymupdf.open(cel) as pdf:
        assert pdf.page_count == len(p.arkusze)
        r = pdf[0].rect
        assert (round(r.width * 25.4 / 72), round(r.height * 25.4 / 72)) == (420, 297)  # A3


@pytest.mark.parametrize("wgs84", [False, True])
def test_pismo(dxf, operator, wgs84):
    p = nowy_projekt(dxf, wgs84=wgs84)
    cel = dxf.with_name("projekt_pismo.docx")
    p.zapisz_pismo(cel, operator, "01.10.2026", "-")
    d = docx.Document(cel)
    tekst = "\n".join(x.text for x in d.paragraphs)
    assert "OD: 01.10.2026 DO: -" in tekst
    assert f"Łączna ilość Słupów elektroenergetycznych: {len(p.plan)}" in tekst
    tabela = next(t for t in d.tables if t.rows[0].cells[0].text.strip().startswith("L.p"))
    assert len(tabela.rows) == len(p.plan) + 1
    x = float(tabela.rows[1].cells[5].text)
    assert (50 < x < 56) if wgs84 else (x > 5_000_000)


def test_zestawienie_rozbudowane(dxf):
    p = nowy_projekt(dxf)
    cel = dxf.with_name("projekt_rozbudowana.xlsx")
    p.zapisz_rozbudowana(cel)
    ws = load_workbook(cel).worksheets[0]
    assert ws.cell(4, 1).value == 1
    assert ws.cell(3 + len(p.plan), 1).value == len(p.plan)
    assert ws.cell(5 + len(p.plan), 1).value == f"Wykaz dotyczy łącznie {len(p.plan)} słupów."
    assert "str. 7" in ws["F3"].value  # EPSG:2178 = strefa 7


def test_zestawienie_uproszczone(dxf):
    p = nowy_projekt(dxf)
    cel = dxf.with_name("projekt_uproszczona.xls")
    p.zapisz_uproszczona(cel)
    s = xlrd.open_workbook(str(cel)).sheet_by_index(0)
    assert s.cell_value(4, 0).strip() == p.plan[0].etykieta
    assert s.cell_value(3 + len(p.plan), 0).strip() == p.plan[-1].etykieta
    assert s.cell_value(4, 5) == pytest.approx(p.plan[0].geo_x, abs=0.01)


def test_tabela_projektowa_zawsze_wgs84(dxf):
    p = nowy_projekt(dxf)  # wgs84=False, a tabela i tak ma GPS
    cel = dxf.with_name("projekt_tabela_projektowa.xlsx")
    p.zapisz_projektowa(cel)
    ws = load_workbook(cel).worksheets[0]
    lat, lon = ws.cell(2, 7).value, ws.cell(2, 8).value
    assert 49 < lat < 55 and 14 < lon < 25
    assert ws.max_row == len(p.plan) + 1


def test_wiersz_polecen(dxf, capsys):
    kat = dxf.parent / "wyniki"
    assert cli([str(dxf), "--bez-geokodowania", "--tak", "--katalog", str(kat),
                "--prefiks", "S", "--format", "dxf"]) == 0
    pliki = sorted(f.name for f in kat.iterdir())
    assert pliki == ["projekt_numeracja.dxf", "projekt_pismo.docx", "projekt_rozbudowana.xlsx",
                     "projekt_tabela_projektowa.xlsx", "projekt_uproszczona.xls"]
    wyjscie = capsys.readouterr().out
    assert "Słupów: 151" in wyjscie and "początek: S1" in wyjscie
